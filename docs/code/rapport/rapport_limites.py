# -*- coding: utf-8 -*-
"""Met a jour la section « Limites identifiees » du rapport.

Deux limites y sont decrites comme subsistantes alors qu'elles ont ete levees
depuis. Les laisser en l'etat ferait dire au rapport le contraire de ce que
dit la section 6.6 quelques paragraphes plus loin, et de ce que montre
l'infrastructure.

Une limite levee ne disparait pas du rapport : elle devient un resultat. La
formulation retenue dit donc ce qui etait, ce qui a ete fait, et ce qui reste.
"""
import os

from docx import Document

DOSSIER = r"C:/Users/AMBER/OneDrive - ST DIGITAL SARL/Documents/Dossier Stage"
FICHIER = os.path.join(DOSSIER, "RapportStageSTD_v6.docx")

REMPLACEMENTS = [
    (
        "La base vectorielle est conservée en mémoire",
        "La base vectorielle était initialement conservée en mémoire, et tout redémarrage du "
        "service d'orchestration la vidait : l'assistant se mettait alors à répondre sans "
        "contexte, sans qu'aucune erreur ne soit émise. Le défaut s'étant produit deux fois, il "
        "a été corrigé en fin de stage par le passage à une base vectorielle persistante, "
        "PostgreSQL doté de l'extension pgvector, hébergée sur la même infrastructure interne. "
        "La correction a été vérifiée par un redémarrage sans réindexation, à l'issue duquel les "
        "réponses sont restées exactes."
    ),
    (
        "La surface publique expose, pendant la phase d'expérimentation",
        "La surface publique exposait, pendant la phase d'expérimentation, davantage que le "
        "strict nécessaire, en particulier l'interface d'administration. Elle a été refermée : "
        "seules l'application visiteur, la conversation et la page de réponse de disponibilité y "
        "demeurent joignables. Une première tentative de correction s'était révélée sans effet, "
        "le serveur mandataire appliquant ses règles par préfixe, ce qui n'a été établi que par "
        "une requête réellement émise et non par relecture de la configuration."
    ),
    (
        "Les comptes d'administration reposent sur un mot de passe partagé",
        "Les comptes d'administration reposent sur un mot de passe partagé par pays, ce qui ne "
        "permet ni d'identifier l'auteur d'une décision, ni de retirer l'accès à une personne en "
        "particulier. À défaut de comptes nominatifs, un journal d'activité a été mis en place en "
        "fin de stage : il enregistre les connexions, les refus de connexion, les changements de "
        "statut de rendez-vous, les clôtures de visite et les changements de mot de passe. Il "
        "indique le site depuis lequel l'action a été prononcée, et non la personne : la limite "
        "est atténuée, elle n'est pas levée."
    ),
]

d = Document(FICHIER)
faits = []
for p in d.paragraphs:
    if p.style.name.startswith("TOC"):
        continue
    t = p.text.strip()
    for debut, nouveau in REMPLACEMENTS:
        if t.startswith(debut):
            # On vide les runs existants et on en repose un seul, pour ne pas
            # heriter d'une mise en forme partielle repartie sur plusieurs runs.
            garde = p.runs[0] if p.runs else None
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            nr = p.add_run(nouveau)
            if garde is not None:
                nr.font.name = garde.font.name
                nr.font.size = garde.font.size
            faits.append(debut[:48])
            break

d.save(FICHIER)
print("limites mises a jour : %d" % len(faits))
for f in faits:
    print("   %s..." % f)
if len(faits) != len(REMPLACEMENTS):
    print("ATTENTION : %d remplacement(s) non trouve(s)" % (len(REMPLACEMENTS) - len(faits)))
